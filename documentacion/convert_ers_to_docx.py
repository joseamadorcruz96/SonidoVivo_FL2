import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Establece el color de fondo de una celda en formato HEX (ej. '002B49')."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Establece padding interno en dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Aplica bordes finos y limpios a toda la tabla."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_styled_paragraph_runs(paragraph, text):
    """Parsea markdown inline como negrita (**texto**), cursiva (*texto*) y código (`codigo`)."""
    pattern = r'(\*\*.*?\*\*|\*.*?\*|`.*?`)'
    parts = re.split(pattern, text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(160, 50, 50)
        else:
            paragraph.add_run(part)

def markdown_to_docx(md_path, docx_path):
    print(f"Leyendo archivo Markdown: {md_path}")
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Configurar márgenes a 2.5 cm en todas las direcciones
    for section in doc.sections:
        section.top_margin = Inches(0.98)
        section.bottom_margin = Inches(0.98)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.98)

    # Configuración de estilos base
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(45, 55, 72)

    # Variables de estado para el parser
    in_table = False
    table_rows = []
    in_code_block = False
    code_lines = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        
        # Limpiar filas separadoras (ej. | :--- | :--- |)
        filtered_rows = []
        for r in table_rows:
            if all(re.match(r'^[\s:-]+$', col) for col in r if col):
                continue
            filtered_rows.append(r)

        if not filtered_rows:
            in_table = False
            table_rows = []
            return

        num_cols = max(len(r) for r in filtered_rows)
        # Normalizar filas para que todas tengan num_cols
        for i in range(len(filtered_rows)):
            while len(filtered_rows[i]) < num_cols:
                filtered_rows[i].append("")

        table = doc.add_table(rows=len(filtered_rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)

        for row_idx, row_data in enumerate(filtered_rows):
            is_header = (row_idx == 0)
            row = table.rows[row_idx]
            
            # Repetir encabezado en cada página
            if is_header:
                trPr = row._tr.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for col_idx, cell_value in enumerate(row_data):
                cell = row.cells[col_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=140, bottom=140, left=160, right=160)

                # Formatear párrafos dentro de la celda
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15

                if is_header:
                    set_cell_background(cell, "002B49")  # Azul Marino Duoc UC
                    lines_in_cell = cell_value.split("<br>")
                    for l_idx, line_item in enumerate(lines_in_cell):
                        if l_idx > 0:
                            p = cell.add_paragraph()
                            p.paragraph_format.space_before = Pt(1)
                            p.paragraph_format.space_after = Pt(1)
                        run = p.add_run(line_item.strip())
                        run.font.name = 'Calibri'
                        run.font.bold = True
                        run.font.size = Pt(10)
                        run.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "F8FAFC")  # Fondo alterno suave
                    else:
                        set_cell_background(cell, "FFFFFF")
                    
                    lines_in_cell = cell_value.split("<br>")
                    for l_idx, line_item in enumerate(lines_in_cell):
                        if l_idx > 0:
                            p = cell.add_paragraph()
                            p.paragraph_format.space_before = Pt(1)
                            p.paragraph_format.space_after = Pt(1)
                        p.paragraph_format.line_spacing = 1.15
                        add_styled_paragraph_runs(p, line_item.strip())
                        for r in p.runs:
                            if not r.bold and not r.italic:
                                r.font.name = 'Calibri'
                                r.font.size = Pt(9.5)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)
        in_table = False
        table_rows = []

    def flush_code():
        nonlocal in_code_block, code_lines
        if code_lines:
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, top=160, bottom=160, left=200, right=200)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for c_line in code_lines:
                run = p.add_run(c_line + "\n")
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(30, 41, 59)
            doc.add_paragraph()
        in_code_block = False
        code_lines = []

    idx = 0
    total_lines = len(lines)
    while idx < total_lines:
        raw_line = lines[idx]
        line = raw_line.rstrip('\r\n')
        stripped = line.strip()

        # Manejo de bloques de código
        if stripped.startswith("```"):
            if in_code_block:
                flush_code()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
                code_lines = []
            idx += 1
            continue

        if in_code_block:
            code_lines.append(line)
            idx += 1
            continue

        # Manejo de tablas
        if stripped.startswith("|") and stripped.endswith("|"):
            in_table = True
            # Extraer celdas
            cells = [c.strip() for c in stripped.split("|")[1:-1]]
            table_rows.append(cells)
            idx += 1
            continue
        elif in_table:
            flush_table()

        # Línea en blanco
        if not stripped:
            idx += 1
            continue

        # Separador horizontal
        if stripped in ["---", "***", "___"]:
            # Línea divisoria o salto
            idx += 1
            continue

        # Encabezados
        if stripped.startswith("# "):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0, 43, 73)  # Azul Institucional Duoc
        elif stripped.startswith("## "):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0, 80, 122)
        elif stripped.startswith("### "):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.bold = True
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(197, 137, 24)  # Oro Luthería
        elif stripped.startswith("#### "):
            text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(31, 41, 55)
        # Listas no ordenadas (viñetas)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            content = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_styled_paragraph_runs(p, content)
        # Listas ordenadas
        elif re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^\d+\.\s', stripped)
            content = stripped[match.end():].strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_styled_paragraph_runs(p, content)
        else:
            # Párrafo normal
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15
            add_styled_paragraph_runs(p, stripped)

        idx += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    print(f"Guardando documento DOCX en: {docx_path}")
    doc.save(docx_path)
    print("Conversión finalizada exitosamente.")

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    md_file = os.path.join(base_dir, "Informe_ERS.md")
    docx_file = os.path.join(base_dir, "Informe_ERS.docx")
    markdown_to_docx(md_file, docx_file)
