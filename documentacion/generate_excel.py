import re
import os
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def clean_cell(text):
    if not text:
        return ""
    text = text.strip()
    # Replace <br> or <br/> or <br /> with newline
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.IGNORECASE)
    # Remove markdown bold/italic markers
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    # Remove backticks
    text = text.replace('`', '')
    return text.strip()

def parse_markdown(md_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    definitions = []
    rf_rows = []
    rnf_rows = []
    
    current_section = None
    
    for line in lines:
        line_str = line.strip()
        if "Definiciones de Columnas y Criterios" in line_str:
            current_section = "defs"
            continue
        elif "1. Matriz de Requerimientos Funcionales" in line_str:
            current_section = "rf"
            continue
        elif "2. Matriz de Requerimientos No Funcionales" in line_str:
            current_section = "rnf"
            continue
            
        if line_str.startswith('|') and line_str.endswith('|'):
            # Skip separator line like | :--- | :--- |
            if re.match(r'^\|[\s:-]+\|', line_str) and not any(c.isalnum() for c in line_str):
                continue
            
            parts = [clean_cell(p) for p in line_str.split('|')[1:-1]]
            
            if not parts or not any(parts):
                continue
                
            # Skip only the header rows
            if current_section == "defs":
                if parts[0].strip().lower() in ("columna", "[columna]"):
                    continue
                if len(parts) >= 2:
                    definitions.append(parts[:2])
            elif current_section in ("rf", "rnf"):
                if parts[0].strip() in ("[R-N°]", "R-N°", "[R-Nº]"):
                    continue
                if len(parts) >= 8:
                    if current_section == "rf":
                        rf_rows.append(parts[:8])
                    else:
                        rnf_rows.append(parts[:8])
                
    return definitions, rf_rows, rnf_rows

def build_excel(md_path, excel_path):
    definitions, rf_rows, rnf_rows = parse_markdown(md_path)
    all_rows = rf_rows + rnf_rows
    
    wb = Workbook()
    
    # -------------------------------------------------------------
    # STYLES DEFINITION
    # -------------------------------------------------------------
    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="1F2937")
    font_bold_data = Font(name="Calibri", size=10, bold=True, color="111827")
    font_status = Font(name="Calibri", size=10, bold=True, color="166534")
    
    # Section Header & KPI Fonts
    font_section_h = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    font_kpi_num = Font(name="Calibri", size=20, bold=True, color="0F172A")
    font_kpi_label = Font(name="Calibri", size=9, bold=True, color="475569")
    
    # Colors & Fills
    fill_navy = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_slate_h = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_blue_h = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    fill_purple_h = PatternFill(start_color="6D28D9", end_color="6D28D9", fill_type="solid")
    fill_teal_h = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
    fill_amber_h = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
    
    fill_zebra_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_zebra_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_status_aprobado = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    
    fill_kpi_bg = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_kpi_green = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    
    # Borders
    border_thin_side = Side(border_style="thin", color="CBD5E1")
    border_thick_side = Side(border_style="medium", color="64748B")
    cell_border = Border(left=border_thin_side, right=border_thin_side, top=border_thin_side, bottom=border_thin_side)
    kpi_border = Border(left=border_thin_side, right=border_thin_side, top=border_thin_side, bottom=border_thin_side)
    
    # Alignments
    align_center_top = Alignment(horizontal="center", vertical="top", wrap_text=True)
    align_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_center_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    
    req_headers = [
        "[R-N°]",
        "[Nombre del Requerimiento]",
        "Tipo Requerimiento",
        "Clasificación",
        "Actores Relacionados",
        "[Descripción corta del requerimiento]",
        "Criterios de Aceptación",
        "Estado"
    ]
    
    col_widths = {
        1: 12,  # R-N°
        2: 28,  # Nombre
        3: 18,  # Tipo
        4: 25,  # Clasificación
        5: 25,  # Actores
        6: 48,  # Descripción
        7: 64,  # Criterios
        8: 14   # Estado
    }
    
    # =============================================================
    # SHEET 1: PORTADA Y RESUMEN EJECUTIVO
    # =============================================================
    ws_cover = wb.active
    ws_cover.title = "Portada y Resumen"
    ws_cover.views.sheetView[0].showGridLines = True
    ws_cover.sheet_properties.tabColor = "0F172A"
    
    # Header Banner
    ws_cover.merge_cells("A1:G1")
    ws_cover["A1"] = "PLANILLA DE ESPECIFICACIÓN DE REQUERIMIENTOS DE SOFTWARE (ERS)"
    ws_cover["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    ws_cover["A1"].fill = fill_navy
    ws_cover["A1"].alignment = align_center_center
    ws_cover.row_dimensions[1].height = 36
    
    ws_cover.merge_cells("A2:G2")
    ws_cover["A2"] = 'Proyecto: Tienda Online "Sonido Vivo" | Evaluación Parcial 1 (Forma B) — DSY1104 Desarrollo Full Stack II'
    ws_cover["A2"].font = font_subtitle
    ws_cover["A2"].fill = fill_navy
    ws_cover["A2"].alignment = align_center_center
    ws_cover.row_dimensions[2].height = 22
    
    ws_cover.row_dimensions[3].height = 12
    
    # KPI CARDS ROW (Row 4 & 5)
    # Card 1: Total Requerimientos (Cols A-B)
    ws_cover.merge_cells("A4:B4")
    ws_cover["A4"] = "TOTAL REQUERIMIENTOS"
    ws_cover["A4"].font = font_kpi_label
    ws_cover["A4"].fill = fill_kpi_bg
    ws_cover["A4"].alignment = align_center_center
    
    ws_cover.merge_cells("A5:B5")
    ws_cover["A5"] = len(all_rows)
    ws_cover["A5"].font = font_kpi_num
    ws_cover["A5"].fill = fill_kpi_bg
    ws_cover["A5"].alignment = align_center_center
    
    # Card 2: Requerimientos Funcionales (Cols C)
    ws_cover["C4"] = "REQ. FUNCIONALES (RF)"
    ws_cover["C4"].font = font_kpi_label
    ws_cover["C4"].fill = fill_kpi_bg
    ws_cover["C4"].alignment = align_center_center
    
    ws_cover["C5"] = len(rf_rows)
    ws_cover["C5"].font = font_kpi_num
    ws_cover["C5"].fill = fill_kpi_bg
    ws_cover["C5"].alignment = align_center_center
    
    # Card 3: Requerimientos No Funcionales (Cols D)
    ws_cover["D4"] = "REQ. NO FUNCIONALES (RNF)"
    ws_cover["D4"].font = font_kpi_label
    ws_cover["D4"].fill = fill_kpi_bg
    ws_cover["D4"].alignment = align_center_center
    
    ws_cover["D5"] = len(rnf_rows)
    ws_cover["D5"].font = font_kpi_num
    ws_cover["D5"].fill = fill_kpi_bg
    ws_cover["D5"].alignment = align_center_center
    
    # Card 4: Estado General (Cols E-G)
    ws_cover.merge_cells("E4:G4")
    ws_cover["E4"] = "ESTADO GENERAL DE REQUERIMIENTOS"
    ws_cover["E4"].font = font_kpi_label
    ws_cover["E4"].fill = fill_kpi_green
    ws_cover["E4"].alignment = align_center_center
    
    ws_cover.merge_cells("E5:G5")
    ws_cover["E5"] = "100% APROBADO (25/25)"
    ws_cover["E5"].font = Font(name="Calibri", size=16, bold=True, color="166534")
    ws_cover["E5"].fill = fill_kpi_green
    ws_cover["E5"].alignment = align_center_center
    
    ws_cover.row_dimensions[4].height = 18
    ws_cover.row_dimensions[5].height = 32
    
    for r in (4, 5):
        for col in range(1, 8):
            ws_cover.cell(row=r, column=col).border = cell_border
            
    ws_cover.row_dimensions[6].height = 14
    
    # Table 1: Desglose por Clasificación (Cols A-D, Rows 7-15)
    ws_cover.merge_cells("A7:D7")
    ws_cover["A7"] = "DESGLOSE POR TIPO Y CLASIFICACIÓN"
    ws_cover["A7"].font = font_section_h
    ws_cover["A7"].fill = fill_slate_h
    ws_cover["A7"].alignment = align_center_center
    ws_cover.row_dimensions[7].height = 24
    
    ws_cover.cell(row=8, column=1, value="Tipo").font = font_header
    ws_cover.cell(row=8, column=1).fill = fill_navy
    ws_cover.cell(row=8, column=1).alignment = align_center_center
    ws_cover.cell(row=8, column=1).border = cell_border
    
    ws_cover.merge_cells("B8:C8")
    ws_cover["B8"] = "Clasificación Formal (Anexo 2)"
    ws_cover["B8"].font = font_header
    ws_cover["B8"].fill = fill_navy
    ws_cover["B8"].alignment = align_center_center
    ws_cover["B8"].border = cell_border
    ws_cover["C8"].border = cell_border
    
    ws_cover.cell(row=8, column=4, value="Cantidad").font = font_header
    ws_cover.cell(row=8, column=4).fill = fill_navy
    ws_cover.cell(row=8, column=4).alignment = align_center_center
    ws_cover.cell(row=8, column=4).border = cell_border
    ws_cover.row_dimensions[8].height = 20
    
    breakdown_data = [
        ("Funcional", "Funcional de sistema", sum(1 for r in rf_rows if r[3] == "Funcional de sistema")),
        ("Funcional", "Funcional de usuario", sum(1 for r in rf_rows if r[3] == "Funcional de usuario")),
        ("No Funcional", "No funcional de producto", sum(1 for r in rnf_rows if r[3] == "No funcional de producto")),
        ("No Funcional", "No funcional de sistema", sum(1 for r in rnf_rows if r[3] == "No funcional de sistema")),
        ("No Funcional", "No funcional de la Organización", sum(1 for r in rnf_rows if "Organización" in r[3] or "Organizacion" in r[3])),
        ("No Funcional", "No funcional Externos", sum(1 for r in rnf_rows if "Externos" in r[3])),
    ]
    
    for idx, (t, c, count) in enumerate(breakdown_data, start=9):
        is_even = (idx % 2 == 0)
        row_fill = fill_zebra_even if is_even else fill_zebra_odd
        ws_cover.row_dimensions[idx].height = 20
        
        c1 = ws_cover.cell(row=idx, column=1, value=t)
        c1.font = font_bold_data
        c1.fill = row_fill
        c1.alignment = align_center_center
        c1.border = cell_border
        
        ws_cover.merge_cells(f"B{idx}:C{idx}")
        c2 = ws_cover.cell(row=idx, column=2, value=c)
        c2.font = font_data
        c2.fill = row_fill
        c2.alignment = Alignment(horizontal="left", vertical="center")
        c2.border = cell_border
        ws_cover.cell(row=idx, column=3).border = cell_border
        
        c4 = ws_cover.cell(row=idx, column=4, value=count)
        c4.font = font_bold_data
        c4.fill = row_fill
        c4.alignment = align_center_center
        c4.border = cell_border
        
    # Total Row Breakdown
    total_r = 9 + len(breakdown_data)
    ws_cover.row_dimensions[total_r].height = 22
    ws_cover.merge_cells(f"A{total_r}:C{total_r}")
    t_label = ws_cover.cell(row=total_r, column=1, value="TOTAL DE REQUERIMIENTOS ESPECIFICADOS")
    t_label.font = font_bold_data
    t_label.fill = fill_kpi_bg
    t_label.alignment = Alignment(horizontal="right", vertical="center")
    t_label.border = cell_border
    ws_cover.cell(row=total_r, column=2).border = cell_border
    ws_cover.cell(row=total_r, column=3).border = cell_border
    
    t_val = ws_cover.cell(row=total_r, column=4, value=len(all_rows))
    t_val.font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    t_val.fill = fill_kpi_bg
    t_val.alignment = align_center_center
    t_val.border = cell_border
    
    # Table 2: Índice de Pestañas y Contenido (Cols E-G, Rows 7-15)
    ws_cover.merge_cells("E7:G7")
    ws_cover["E7"] = "GUÍA DE NAVEGACIÓN DEL LIBRO"
    ws_cover["E7"].font = font_section_h
    ws_cover["E7"].fill = fill_slate_h
    ws_cover["E7"].alignment = align_center_center
    
    ws_cover.cell(row=8, column=5, value="Pestaña").font = font_header
    ws_cover.cell(row=8, column=5).fill = fill_navy
    ws_cover.cell(row=8, column=5).alignment = align_center_center
    ws_cover.cell(row=8, column=5).border = cell_border
    
    ws_cover.merge_cells("F8:G8")
    ws_cover["F8"] = "Descripción del Contenido"
    ws_cover["F8"].font = font_header
    ws_cover["F8"].fill = fill_navy
    ws_cover["F8"].alignment = align_center_center
    ws_cover["F8"].border = cell_border
    ws_cover["G8"].border = cell_border
    
    tabs_guide = [
        ("Req. Funcionales (RF)", "Matriz completa con los 17 Requerimientos Funcionales (RF-01 a RF-17)."),
        ("Req. No Funcionales (RNF)", "Matriz completa con los 8 Requerimientos No Funcionales (RNF-01 a RNF-08)."),
        ("Matriz Consolidada (RF + RNF)", "Vista unificada con los 25 requerimientos del proyecto."),
        ("Definiciones (Anexo 2)", "Definiciones oficiales de columnas y criterios según pauta Anexo 2."),
    ]
    
    for idx, (t_name, t_desc) in enumerate(tabs_guide, start=9):
        is_even = (idx % 2 == 0)
        row_fill = fill_zebra_even if is_even else fill_zebra_odd
        ws_cover.row_dimensions[idx].height = 22
        
        c5 = ws_cover.cell(row=idx, column=5, value=t_name)
        c5.font = font_bold_data
        c5.fill = row_fill
        c5.alignment = Alignment(horizontal="center", vertical="center")
        c5.border = cell_border
        
        ws_cover.merge_cells(f"F{idx}:G{idx}")
        c6 = ws_cover.cell(row=idx, column=6, value=t_desc)
        c6.font = font_data
        c6.fill = row_fill
        c6.alignment = Alignment(horizontal="left", vertical="center")
        c6.border = cell_border
        ws_cover.cell(row=idx, column=7).border = cell_border
        
    # Fill remaining empty row on right side
    empty_r = total_r
    ws_cover.cell(row=empty_r, column=5).border = cell_border
    ws_cover.cell(row=empty_r, column=5).fill = fill_kpi_bg
    ws_cover.merge_cells(f"F{empty_r}:G{empty_r}")
    ws_cover.cell(row=empty_r, column=6).border = cell_border
    ws_cover.cell(row=empty_r, column=7).border = cell_border
    ws_cover.cell(row=empty_r, column=6).fill = fill_kpi_bg
    ws_cover.cell(row=empty_r, column=6, value="Cumplimiento estricto de pauta Evaluación Parcial 1").font = font_subtitle
    ws_cover.cell(row=empty_r, column=6).alignment = align_center_center
    
    ws_cover.column_dimensions["A"].width = 16
    ws_cover.column_dimensions["B"].width = 24
    ws_cover.column_dimensions["C"].width = 24
    ws_cover.column_dimensions["D"].width = 16
    ws_cover.column_dimensions["E"].width = 28
    ws_cover.column_dimensions["F"].width = 38
    ws_cover.column_dimensions["G"].width = 38

    # =============================================================
    # HELPER TO POPULATE REQUIREMENT MATRICES
    # =============================================================
    def populate_req_sheet(ws, title, subtitle, rows, header_color, tab_color):
        ws.views.sheetView[0].showGridLines = True
        ws.sheet_properties.tabColor = tab_color
        
        # Row 1: Title Banner
        ws.merge_cells("A1:H1")
        ws["A1"] = f"SONIDO VIVO — {title.upper()}"
        ws["A1"].font = font_title
        ws["A1"].fill = fill_navy
        ws["A1"].alignment = align_center_center
        ws.row_dimensions[1].height = 32
        
        # Row 2: Subtitle
        ws.merge_cells("A2:H2")
        ws["A2"] = f"Proyecto: Tienda Online Sonido Vivo | DSY1104 Desarrollo Full Stack II | {subtitle}"
        ws["A2"].font = font_subtitle
        ws["A2"].fill = fill_navy
        ws["A2"].alignment = align_center_center
        ws.row_dimensions[2].height = 20
        
        # Row 3: Empty spacing
        ws.row_dimensions[3].height = 8
        
        # Row 4: Headers
        ws.row_dimensions[4].height = 28
        fill_col_h = PatternFill(start_color=header_color, end_color=header_color, fill_type="solid")
        for col_idx, h_text in enumerate(req_headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=h_text)
            cell.font = font_header
            cell.fill = fill_col_h
            cell.alignment = align_center_center
            cell.border = cell_border
            
        # Data Rows
        start_row = 5
        for r_idx, row_data in enumerate(rows, start=start_row):
            is_even = (r_idx % 2 == 0)
            row_fill = fill_zebra_even if is_even else fill_zebra_odd
            
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                cell.border = cell_border
                cell.fill = row_fill
                
                # Column specific styling
                if c_idx == 1: # R-N°
                    cell.font = font_bold_data
                    cell.alignment = align_center_top
                elif c_idx == 2: # Nombre
                    cell.font = font_bold_data
                    cell.alignment = align_left_top
                elif c_idx in (3, 4): # Tipo, Clasificación
                    cell.font = font_data
                    cell.alignment = align_center_top if c_idx == 3 else align_left_top
                elif c_idx == 5: # Actores
                    cell.font = font_data
                    cell.alignment = align_left_top
                elif c_idx in (6, 7): # Descripción, Criterios
                    cell.font = font_data
                    cell.alignment = align_left_top
                elif c_idx == 8: # Estado
                    cell.font = font_status
                    cell.alignment = align_center_top
                    cell.fill = fill_status_aprobado
                    
        # Apply Column Widths
        for c_idx, width in col_widths.items():
            ws.column_dimensions[get_column_letter(c_idx)].width = width
            
        # Auto-filter
        last_row = start_row + len(rows) - 1
        ws.auto_filter.ref = f"A4:H{last_row}"
        
        # Freeze panes at row 5 so header stays visible while scrolling
        ws.freeze_panes = "A5"

    # =============================================================
    # SHEET 2: REQUERIMIENTOS FUNCIONALES (RF)
    # =============================================================
    ws_rf = wb.create_sheet(title="Req. Funcionales (RF)")
    populate_req_sheet(
        ws_rf,
        "Matriz de Requerimientos Funcionales",
        "Total: 17 Requerimientos Funcionales (RF-01 al RF-17)",
        rf_rows,
        header_color="1E3A8A", # Darker blue
        tab_color="2563EB"     # Blue
    )
    
    # =============================================================
    # SHEET 3: REQUERIMIENTOS NO FUNCIONALES (RNF)
    # =============================================================
    ws_rnf = wb.create_sheet(title="Req. No Funcionales (RNF)")
    populate_req_sheet(
        ws_rnf,
        "Matriz de Requerimientos No Funcionales",
        "Total: 8 Requerimientos No Funcionales (RNF-01 al RNF-08)",
        rnf_rows,
        header_color="581C87", # Dark purple
        tab_color="7C3AED"     # Purple
    )
    
    # =============================================================
    # SHEET 4: MATRIZ CONSOLIDADA (RF + RNF)
    # =============================================================
    ws_all = wb.create_sheet(title="Matriz Consolidada (RF + RNF)")
    populate_req_sheet(
        ws_all,
        "Matriz Consolidada de Requerimientos",
        "Total: 25 Requerimientos (17 RF + 8 RNF)",
        all_rows,
        header_color="115E59", # Dark teal
        tab_color="0D9488"     # Teal
    )
    
    # =============================================================
    # SHEET 5: DEFINICIONES (ANEXO 2)
    # =============================================================
    ws_def = wb.create_sheet(title="Definiciones (Anexo 2)")
    ws_def.views.sheetView[0].showGridLines = True
    ws_def.sheet_properties.tabColor = "D97706" # Amber
    
    ws_def.merge_cells("A1:B1")
    ws_def["A1"] = "DEFINICIONES DE COLUMNAS Y CRITERIOS (ANEXO 2)"
    ws_def["A1"].font = font_title
    ws_def["A1"].fill = fill_navy
    ws_def["A1"].alignment = align_center_center
    ws_def.row_dimensions[1].height = 32
    
    ws_def.merge_cells("A2:B2")
    ws_def["A2"] = "Estándar establecido en la Pauta de Evaluación DSY1104 (Evaluación Parcial 1)"
    ws_def["A2"].font = font_subtitle
    ws_def["A2"].fill = fill_navy
    ws_def["A2"].alignment = align_center_center
    ws_def.row_dimensions[2].height = 20
    
    ws_def.row_dimensions[3].height = 8
    
    ws_def.cell(row=4, column=1, value="Columna / Campo").font = font_header
    ws_def.cell(row=4, column=1).fill = PatternFill(start_color="92400E", end_color="92400E", fill_type="solid")
    ws_def.cell(row=4, column=1).alignment = align_center_center
    ws_def.cell(row=4, column=1).border = cell_border
    
    ws_def.cell(row=4, column=2, value="Definición e Instrucción de Llenado").font = font_header
    ws_def.cell(row=4, column=2).fill = PatternFill(start_color="92400E", end_color="92400E", fill_type="solid")
    ws_def.cell(row=4, column=2).alignment = align_center_center
    ws_def.cell(row=4, column=2).border = cell_border
    ws_def.row_dimensions[4].height = 26
    
    for idx, (col_name, col_desc) in enumerate(definitions, start=5):
        is_even = (idx % 2 == 0)
        row_fill = fill_zebra_even if is_even else fill_zebra_odd
        ws_def.row_dimensions[idx].height = 26
        
        c1 = ws_def.cell(row=idx, column=1, value=col_name)
        c1.font = font_bold_data
        c1.alignment = Alignment(horizontal="left", vertical="center")
        c1.fill = row_fill
        c1.border = cell_border
        
        c2 = ws_def.cell(row=idx, column=2, value=col_desc)
        c2.font = font_data
        c2.alignment = Alignment(horizontal="left", vertical="center")
        c2.fill = row_fill
        c2.border = cell_border
        
    ws_def.column_dimensions["A"].width = 32
    ws_def.column_dimensions["B"].width = 95
    ws_def.freeze_panes = "A5"
    
    wb.save(excel_path)
    print(f"Excel generado exitosamente en: {excel_path}")
    print(f"Hojas: {wb.sheetnames}")
    print(f"Total RF: {len(rf_rows)}, Total RNF: {len(rnf_rows)}, Total Definiciones: {len(definitions)}")

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    md_file = os.path.join(base_dir, "Requerimientos_Funcionales_simplificado_1.md")
    excel_file = os.path.join(base_dir, "Requerimientos_Funcionales_simplificado_1.xlsx")
    build_excel(md_file, excel_file)
